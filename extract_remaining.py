from pathlib import Path
# pyrefly: ignore [missing-import]
from pypdf import PdfReader
# pyrefly: ignore [missing-import]
from docx import Document

RAW_DIR = Path(r"D:\Computer Science\NLCSNKHMT\chatbot-ctu-NLCSNKHMT\Data_CTU_restructured\raw")
OUTPUT_DIR = Path(r"D:\Computer Science\NLCSNKHMT\chatbot-ctu-NLCSNKHMT\Data_CTU_restructured\processed")

# Các file đã có lớp text tốt -> trích xuất thẳng, không cần OCR
PDF_TARGETS = [
    "ktx/Thong_bao_dang_ky_KTX_HK1_2026_2027.pdf",
    "hoc_bong_ctsv/Cong_tac_sinh_vien.pdf",
    "dang_ky_hoc_phan/Ke_hoach_dang_ky_hoc_phan_HK1_2026_2027.pdf",
    "handbook/So_tay_sinh_vien_2025.pdf",
]

DOCX_TARGETS = [
    "danh_muc_nganh/danh_muc_nganh.docx",
]

# CỐ Ý KHÔNG xử lý:
# ktx/DS_khong_du_dieu_kien_KTX.pdf -> chứa danh sách MSSV/họ tên sinh viên,
# không đưa vào kho tri thức RAG vì lý do bảo vệ dữ liệu cá nhân.


def extract_pdf(pdf_path: Path, out_txt: Path):
    reader = PdfReader(str(pdf_path))
    text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    out_txt.write_text(text, encoding="utf-8")


def extract_docx(docx_path: Path, out_txt: Path):
    doc = Document(str(docx_path))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    out_txt.write_text("\n".join(lines), encoding="utf-8")


def run():
    for rel_path in PDF_TARGETS:
        src = RAW_DIR / rel_path
        if not src.exists():
            print(f"Cảnh báo: không tìm thấy {src}")
            continue
        out_path = (OUTPUT_DIR / rel_path).with_suffix(".md")
        out_path.parent.mkdir(parents=True, exist_ok=True)
        print(f"[PDF] {src.name} -> {out_path}")
        try:
            extract_pdf(src, out_path)
            n = len(out_path.read_text(encoding="utf-8"))
            print(f"  -> xong ({n} ký tự)")
        except Exception as e:
            print(f"  -> LỖI: {e}")

    for rel_path in DOCX_TARGETS:
        src = RAW_DIR / rel_path
        if not src.exists():
            print(f"Cảnh báo: không tìm thấy {src}")
            continue
        out_path = (OUTPUT_DIR / rel_path).with_suffix(".md")
        out_path.parent.mkdir(parents=True, exist_ok=True)
        print(f"[DOCX] {src.name} -> {out_path}")
        try:
            extract_docx(src, out_path)
            n = len(out_path.read_text(encoding="utf-8"))
            print(f"  -> xong ({n} ký tự)")
        except Exception as e:
            print(f"  -> LỖI: {e}")


if __name__ == "__main__":
    run()